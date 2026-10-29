from __future__ import annotations

import io
from typing import Optional

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    buf = io.StringIO()
    df.to_csv(buf, index=False, encoding="utf-8-sig")
    return buf.getvalue().encode("utf-8-sig")


def to_excel_bytes(
    df_jadwal:  pd.DataFrame,
    tabel_hs:   pd.DataFrame,
    df_runs:    pd.DataFrame,
    ringkasan:  pd.DataFrame,
) -> bytes:
    buf = io.BytesIO()

    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        df_jadwal.to_excel(xw, sheet_name="Daftar Jadwal",     index=False)
        tabel_hs.to_excel(xw,  sheet_name="Hari x Slot")
        df_runs.to_excel(xw,   sheet_name="Hasil 10 Run",      index=False)
        ringkasan.to_excel(xw, sheet_name="Ringkasan 10 Run")

        _format_workbook(xw)

    return buf.getvalue()



def _format_workbook(xw: pd.ExcelWriter) -> None:
    header_font  = Font(bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", fgColor="305496")

    for ws in xw.book.worksheets:
        # Header biru di baris pertama
        for cell in ws[1]:
            cell.font  = header_font
            cell.fill  = header_fill

        # Lebar kolom otomatis berdasarkan baris terpanjang
        for col in ws.columns:
            lebar = max(
                (len(str(cell.value).split("\n")[0]) if cell.value else 0)
                for cell in col
            )
            ws.column_dimensions[col[0].column_letter].width = min(max(12, lebar + 2), 42)

    # Sheet 'Hari x Slot' — tinggi baris dan wrap text
    ws_hs = xw.book["Hari x Slot"]
    for row in ws_hs.iter_rows(min_row=2):
        ws_hs.row_dimensions[row[0].row].height = 110
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
