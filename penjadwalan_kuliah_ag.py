import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

try:
    from IPython import get_ipython
    if get_ipython() is not None:
        from IPython.display import display
    else:
        def display(*args):
            for a in args:
                if hasattr(a, "data"):
                    print(a.data.to_string(index=False))
                elif hasattr(a, "to_string"):
                    print(a.to_string())
                else:
                    print(a)
except ImportError:
    def display(*args):
        for a in args:
            if hasattr(a, "data"):
                print(a.data.to_string(index=False))
            elif hasattr(a, "to_string"):
                print(a.to_string())
            else:
                print(a)

import random
import time
from collections import Counter, defaultdict
from dataclasses import dataclass, replace

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt

pd.set_option("display.max_colwidth", 80)
pd.set_option("display.width", 160)
plt.rcParams.update({"figure.dpi": 110, "axes.grid": True, "grid.alpha": 0.3})

# Seed yang dipakai untuk 10 run pengujian awal (sesuai ketentuan: minimal 10 seed berbeda)
SEEDS = [11, 22, 33, 44, 55, 66, 77, 88, 99, 110]

# ---------- Hari & slot waktu ----------
HARI = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat"]
SLOT = ["07.00–09.30", "09.30–12.00", "12.00–14.30", "14.30–17.00", "17.00–19.30"]
N_HARI, N_SLOT = len(HARI), len(SLOT)
N_TIMESLOT = N_HARI * N_SLOT          # 25 timeslot; t -> hari = t // N_SLOT, slot = t % N_SLOT

def nama_timeslot(t):
    return HARI[t // N_SLOT], SLOT[t % N_SLOT]
# ---------- Dosen: id -> (nama lengkap, nama singkat, slot-of-day yang disukai [0..4]) ----------
DOSEN = {
    "D1": ("Felicia Sylviana, S.T., M.M.",                  "Felicia",  {0, 1, 2}),
    "D2": ("Widiatry, S.T., M.T.",                           "Widiatry", {0, 1}),
    "D3": ("Efrans Christian, S.T., M.T.",                   "Efrans",   {1, 2, 3}),
    "D4": ("Deddy Ronaldo, S.T., M.T.",                      "Deddy",    {0, 1, 2}),
    "D5": ("Devi Karolita, S.Kom., M.Kom., Ph.D.",           "Devi",     {0, 1, 2}),
    "D6": ("Agus Sehatman Saragih, S.T., M.Eng.",            "Agus",     {1, 2, 3}),
    "D7": ("Septian Geges, S.Kom., M.Kom.",                  "Septian",  {0, 1, 2, 3}),
    "D8": ("Nahumi Nugrahaningsih, S.T., M.T., Ph.D.",       "Nahumi",   {0, 1}),
}

# ---------- Ruang: id -> (nama ruang, kapasitas [asumsi]) ----------
RUANG = {
    "R1": ("GKTA-LT4", 50),
    "R2": ("GKTA-LT5", 50),
    "R3": ("FT4",      40),
    "R4": ("FT5",      40),
    "R5": ("FT6",      35),
}

# ---------- Kelas mahasiswa (semester-kelas) -> jumlah mahasiswa [asumsi] ----------
KELAS = {"I-A": 45, "I-B": 38, "III-A": 36, "III-B": 33, "V-A": 30, "V-B": 28}

# ---------- Sesi kuliah: (kode, mata kuliah, kelas, id dosen, SKS) ----------
MATKUL = [
    ("MK01", "Aljabar Linear Matriks",         "I-A",   "D1", 3),
    ("MK02", "Pengantar Teknik Informatika",   "I-A",   "D1", 3),
    ("MK03", "Metodologi Penelitian",          "V-A",   "D1", 3),
    ("MK04", "Logika Diskrit",                 "I-A",   "D2", 3),
    ("MK05", "Algoritma dan Pemrograman I",    "I-A",   "D2", 3),
    ("MK06", "Pemrograman Berorientasi Objek", "III-A", "D2", 3),
    ("MK07", "Logika Diskrit",                 "I-B",   "D3", 3),
    ("MK08", "Pemrograman Berorientasi Objek", "III-B", "D3", 3),
    ("MK09", "Basis Data I",                   "III-B", "D3", 3),
    ("MK10", "Data Mining",                    "V-B",   "D3", 3),
    ("MK11", "Analisis dan Desain Perangkat Lunak", "V-B", "D4", 3),
    ("MK12", "Pengolahan Citra Digital",       "V-B",   "D8", 3),
    ("MK13", "Metodologi Penelitian",          "V-B",   "D5", 3),
    ("MK14", "Interaksi Manusia dan Komputer", "III-A", "D5", 3),
    ("MK15", "Basis Data I",                   "III-A", "D5", 3),
    ("MK16", "Keamanan Data dan Informasi",    "V-A",   "D6", 3),
    ("MK17", "Kompleksitas Algoritma",         "III-A", "D6", 3),
    ("MK18", "Sistem Digital",                 "I-B",   "D7", 3),
    ("MK19", "Jaringan Syaraf Tiruan",         "V-A",   "D7", 3),
    ("MK20", "Sistem Digital",                 "I-A",   "D8", 3),
]

# - Struktur turunan (indeks numerik untuk perhitungan cepat) -
ID_DOSEN = list(DOSEN)
ID_RUANG = list(RUANG)
ID_KELAS = list(KELAS)

N_GEN = len(MATKUL)                   # 1 gen = 1 sesi kuliah
N_RUANG = len(RUANG)

DOSEN_IDX = [ID_DOSEN.index(m[3]) for m in MATKUL]
KELAS_IDX = [ID_KELAS.index(m[2]) for m in MATKUL]
UKURAN    = [KELAS[m[2]] for m in MATKUL]
KAPASITAS = [RUANG[r][1] for r in ID_RUANG]
PREF_SLOT = [DOSEN[d][2] for d in ID_DOSEN]

df_matkul = pd.DataFrame(MATKUL, columns=["Kode", "Mata Kuliah", "Kelas", "ID Dosen", "SKS"])
df_matkul["Dosen"] = df_matkul["ID Dosen"].map(lambda d: DOSEN[d][0])
df_matkul["Jml Mhs"] = UKURAN
print(f"Sesi kuliah: {N_GEN} | Dosen: {len(DOSEN)} | Ruang: {N_RUANG} | Timeslot: {N_HARI} hari x {N_SLOT} slot = {N_TIMESLOT}")
df_matkul

# Ringkasan beban dosen, beban kelas, dan ruang
beban_dosen = df_matkul.groupby("Dosen").size().rename("Jumlah Sesi").to_frame()
beban_kelas = df_matkul.groupby("Kelas").size().rename("Jumlah Sesi").to_frame()
df_ruang = pd.DataFrame([(v[0], v[1]) for v in RUANG.values()], columns=["Ruang", "Kapasitas"])
display(beban_dosen, beban_kelas, df_ruang)

def hitung_pelanggaran(k):
    c_dosen, c_ruang, c_kelas = Counter(), Counter(), Counter()
    kapasitas_kurang = 0
    for i, (t, r) in enumerate(k):
        c_dosen[(t, DOSEN_IDX[i])] += 1
        c_ruang[(t, r)] += 1
        c_kelas[(t, KELAS_IDX[i])] += 1
        if KAPASITAS[r] < UKURAN[i]:
            kapasitas_kurang += 1
    bentrok = lambda c: sum(v - 1 for v in c.values())
    return {"H1_dosen": bentrok(c_dosen), "H2_ruang": bentrok(c_ruang),
            "H3_kelas": bentrok(c_kelas), "H4_kapasitas": kapasitas_kurang}

def total_pelanggaran(k):
    return sum(hitung_pelanggaran(k).values())

def hitung_fitness(k):
    return 1.0 / (1.0 + total_pelanggaran(k))

def hitung_soft(k):
    pref = sum(1 for i, (t, r) in enumerate(k) if (t % N_SLOT) not in PREF_SLOT[DOSEN_IDX[i]])
    per_kelas_hari = defaultdict(list)
    for i, (t, r) in enumerate(k):
        per_kelas_hari[(KELAS_IDX[i], t // N_SLOT)].append(t % N_SLOT)
    jeda = 0
    for slots in per_kelas_hari.values():
        s = sorted(slots)
        jeda += sum(max(0, b - a - 1) for a, b in zip(s, s[1:]))
    return {"S1_preferensi_dosen": pref, "S2_jeda_panjang": jeda}

# - Uji unit fungsi fitness -
# (a) Jadwal sengaja bentrok: semua sesi di timeslot 0 & ruang 0 -> harus ada pelanggaran
k_bentrok = [(0, 0)] * N_GEN
assert total_pelanggaran(k_bentrok) > 0 and hitung_fitness(k_bentrok) < 1.0

# (b) Jadwal valid buatan tangan: tiap sesi di timeslot berbeda, ruang terbesar (kapasitas 50)
k_valid = [(i, 0) for i in range(N_GEN)]
assert total_pelanggaran(k_valid) == 0 and hitung_fitness(k_valid) == 1.0

# (c) Bentrok dosen & kelas: sesi 0 & 1 (dosen D1 sama, kelas I-A sama) di timeslot sama, ruang berbeda
k = list(k_valid); k[0] = (22, 0); k[1] = (22, 1)   # timeslot 22 kosong di k_valid
p = hitung_pelanggaran(k)
assert p["H1_dosen"] == 1 and p["H2_ruang"] == 0 and p["H3_kelas"] == 1   # MK01 & MK02 sama dosen D1 dan kelas I-A

# (d) Bentrok ruang saja: sesi 0 (D1, I-A) & sesi 6 (D3, I-B) di timeslot dan ruang sama
k = list(k_valid); k[0] = (24, 0); k[6] = (24, 0)
p = hitung_pelanggaran(k)
assert p["H2_ruang"] == 1 and p["H1_dosen"] == 0 and p["H3_kelas"] == 0

# (e) Pelanggaran kapasitas: kelas I-A (45 mhs) di ruang FT6 (kapasitas 35)
k = list(k_valid); k[0] = (0, 4)
assert hitung_pelanggaran(k)["H4_kapasitas"] == 1
print("Semua uji unit fitness lulus ✔")

@dataclass
class ParameterGA:
    ukuran_populasi: int = 100
    n_generasi: int = 300
    p_crossover: float = 0.8
    p_mutasi: float = 0.05          # peluang mutasi PER GEN
    metode_seleksi: str = "turnamen"     # "turnamen" | "roulette"
    metode_crossover: str = "uniform"    # "uniform"  | "satu_titik"
    k_turnamen: int = 3
    n_elit: int = 2

# ---------- Inisialisasi ----------
def buat_individu(rng):
    return [(rng.randrange(N_TIMESLOT), rng.randrange(N_RUANG)) for _ in range(N_GEN)]

def buat_populasi(ukuran, rng):
    return [buat_individu(rng) for _ in range(ukuran)]

# ---------- Seleksi ----------
def seleksi_roulette(pop, fit, rng, n):
    # peluang terpilih proporsional dengan fitness
    return rng.choices(pop, weights=fit, k=n)

def seleksi_turnamen(pop, fit, rng, n, k):
    hasil = []
    for _ in range(n):
        peserta = rng.sample(range(len(pop)), k)
        hasil.append(pop[max(peserta, key=lambda i: fit[i])])
    return hasil

def pilih_orangtua(pop, fit, rng, param, n=2):
    if param.metode_seleksi == "roulette":
        return seleksi_roulette(pop, fit, rng, n)
    if param.metode_seleksi == "turnamen":
        return seleksi_turnamen(pop, fit, rng, n, param.k_turnamen)
    raise ValueError("metode_seleksi tidak dikenal")

# ---------- Crossover ----------
def crossover_satu_titik(a, b, rng):
    titik = rng.randint(1, N_GEN - 1)
    return a[:titik] + b[titik:], b[:titik] + a[titik:]

def crossover_uniform(a, b, rng):
    c1, c2 = [], []
    for ga, gb in zip(a, b):
        if rng.random() < 0.5:
            c1.append(ga); c2.append(gb)
        else:
            c1.append(gb); c2.append(ga)
    return c1, c2

def lakukan_crossover(a, b, rng, param):
    if rng.random() >= param.p_crossover:
        return list(a), list(b)                     # tanpa crossover: salin orang tua
    if param.metode_crossover == "satu_titik":
        return crossover_satu_titik(a, b, rng)
    if param.metode_crossover == "uniform":
        return crossover_uniform(a, b, rng)
    raise ValueError("metode_crossover tidak dikenal")

# ---------- Mutasi acak ----------
def mutasi(k, p_mutasi, rng):
    hasil = list(k)
    for i, (t, r) in enumerate(hasil):
        if rng.random() < p_mutasi:
            jenis = rng.choice(("waktu", "ruang", "keduanya"))
            if jenis in ("waktu", "keduanya"):
                t = rng.randrange(N_TIMESLOT)
            if jenis in ("ruang", "keduanya"):
                r = rng.randrange(N_RUANG)
            hasil[i] = (t, r)
    return hasil

def jalankan_ga(param, seed, berhenti_jika_valid=False):
    rng = random.Random(seed)                       # seed terpisah per run -> reprodusibel
    t0 = time.perf_counter()

    pop = buat_populasi(param.ukuran_populasi, rng)
    riwayat_terbaik, riwayat_rata = [], []
    gen_valid = None
    terbaik, fit_terbaik = None, -1.0

    for gen in range(param.n_generasi + 1):
        fit = [hitung_fitness(k) for k in pop]
        urut = sorted(range(len(pop)), key=fit.__getitem__, reverse=True)

        if fit[urut[0]] > fit_terbaik:
            fit_terbaik, terbaik = fit[urut[0]], pop[urut[0]]
        riwayat_terbaik.append(fit[urut[0]])
        riwayat_rata.append(sum(fit) / len(fit))
        if gen_valid is None and fit[urut[0]] == 1.0:
            gen_valid = gen

        if gen == param.n_generasi or (berhenti_jika_valid and gen_valid is not None):
            break

        # --- generasi berikutnya: elitisme + seleksi + crossover + mutasi ---
        baru = [pop[i] for i in urut[:param.n_elit]]
        while len(baru) < param.ukuran_populasi:
            a, b = pilih_orangtua(pop, fit, rng, param, 2)
            c1, c2 = lakukan_crossover(a, b, rng, param)
            baru.append(mutasi(c1, param.p_mutasi, rng))
            if len(baru) < param.ukuran_populasi:
                baru.append(mutasi(c2, param.p_mutasi, rng))
        pop = baru

    return {
        "seed": seed, "kromosom": terbaik, "fitness_terbaik": fit_terbaik,
        "fitness_rata_akhir": riwayat_rata[-1], "gen_valid": gen_valid,
        "riwayat_terbaik": riwayat_terbaik, "riwayat_rata": riwayat_rata,
        "waktu_detik": time.perf_counter() - t0,
    }

# Uji singkat satu run
hasil_uji = jalankan_ga(ParameterGA(), seed=1)
print("Fitness terbaik :", hasil_uji["fitness_terbaik"])
print("Generasi valid  :", hasil_uji["gen_valid"])
print("Waktu           :", round(hasil_uji["waktu_detik"], 2), "detik")

PARAM_UTAMA = ParameterGA(ukuran_populasi=100, n_generasi=300, p_crossover=0.8, p_mutasi=0.05,
                          metode_seleksi="turnamen", metode_crossover="uniform", n_elit=2)

hasil_runs = [jalankan_ga(PARAM_UTAMA, s) for s in SEEDS]

df_runs = pd.DataFrame([{
    "Run": i + 1, "Seed": h["seed"],
    "Fitness Terbaik": h["fitness_terbaik"],
    "Fitness Rata-rata (generasi akhir)": h["fitness_rata_akhir"],
    "Generasi Valid Pertama": h["gen_valid"],
    "Pelanggaran Keras": total_pelanggaran(h["kromosom"]),
    "Waktu (s)": h["waktu_detik"],
} for i, h in enumerate(hasil_runs)])

display(df_runs.round(4).style.hide(axis="index"))

ringkasan = pd.DataFrame({
    "Rata-rata": df_runs.drop(columns=["Run", "Seed"]).mean(numeric_only=True),
    "Simpangan Baku": df_runs.drop(columns=["Run", "Seed"]).std(numeric_only=True),
    "Min": df_runs.drop(columns=["Run", "Seed"]).min(numeric_only=True),
    "Maks": df_runs.drop(columns=["Run", "Seed"]).max(numeric_only=True),
}).round(4)
print(f"Run yang menemukan solusi valid: {df_runs['Generasi Valid Pertama'].notna().sum()} dari {len(df_runs)}")
ringkasan

kombinasi = [(s, c) for s in ("roulette", "turnamen") for c in ("satu_titik", "uniform")]
hasil_kombinasi = {}
for s, c in kombinasi:
    p = replace(PARAM_UTAMA, metode_seleksi=s, metode_crossover=c)
    hasil_kombinasi[(s, c)] = [jalankan_ga(p, sd) for sd in SEEDS]

baris = []
for (s, c), lst in hasil_kombinasi.items():
    gv = [h["gen_valid"] for h in lst if h["gen_valid"] is not None]
    baris.append({
        "Seleksi": s, "Crossover": c,
        "Fitness Terbaik (rata-rata)": np.mean([h["fitness_terbaik"] for h in lst]),
        "Run Valid": f"{len(gv)}/{len(lst)}",
        "Gen Valid (rata-rata)": np.mean(gv) if gv else np.nan,
        "Gen Valid (std)": np.std(gv, ddof=1) if len(gv) > 1 else np.nan,
        "Waktu (s, rata-rata)": np.mean([h["waktu_detik"] for h in lst]),
    })
df_kombinasi = pd.DataFrame(baris)
display(df_kombinasi.round(3).style.hide(axis="index"))

fig, ax = plt.subplots(figsize=(8, 4.2))
for (s, c), lst in hasil_kombinasi.items():
    rata = np.mean([h["riwayat_terbaik"] for h in lst], axis=0)
    ax.plot(rata, label=f"{s} + {c}", linewidth=1.8)
ax.set_xlim(0, 120)
ax.set_xlabel("Generasi"); ax.set_ylabel("Fitness terbaik (rata-rata 10 run)")
ax.set_title("Perbandingan kombinasi seleksi dan crossover (120 generasi pertama)")
ax.legend(); plt.tight_layout(); plt.show()

def kunci_terbaik(h):
    soft = hitung_soft(h["kromosom"])
    return (-h["fitness_terbaik"], sum(soft.values()))

run_terbaik = min(hasil_runs, key=kunci_terbaik)
k_terbaik = run_terbaik["kromosom"]
print(f"Run terbaik: seed {run_terbaik['seed']} | fitness = {run_terbaik['fitness_terbaik']} "
      f"| valid pada generasi {run_terbaik['gen_valid']}")
print("Pelanggaran keras:", hitung_pelanggaran(k_terbaik))
print("Kendala lunak    :", hitung_soft(k_terbaik))

# ---------- Tabel daftar sesi ----------
def bangun_daftar(k):
    baris = []
    for i, (t, r) in enumerate(k):
        kode, mk, kelas, d, sks = MATKUL[i]
        hari, jam = nama_timeslot(t)
        baris.append({"Hari": hari, "Jam": jam, "Kode": kode, "Mata Kuliah": mk, "Kelas": kelas,
                      "Dosen": DOSEN[d][0], "Ruang": RUANG[ID_RUANG[r]][0], "SKS": sks,
                      "_h": t // N_SLOT, "_s": t % N_SLOT})
    df = pd.DataFrame(baris).sort_values(["_h", "_s", "Ruang"]).drop(columns=["_h", "_s"])
    return df.reset_index(drop=True)

df_jadwal = bangun_daftar(k_terbaik)
df_jadwal

# ---------- Tabel jadwal hari x slot ----------
def bangun_tabel_hari_slot(k):
    sel = defaultdict(list)
    for i, (t, r) in enumerate(k):
        kode, mk, kelas, d, sks = MATKUL[i]
        sel[(t % N_SLOT, t // N_SLOT)].append(f"{mk} [{kelas}]\n{DOSEN[d][1]} | {RUANG[ID_RUANG[r]][0]}")
    tabel = pd.DataFrame("", index=SLOT, columns=HARI)
    for (s, h), isi in sel.items():
        tabel.iloc[s, h] = "\n\n".join(isi)
    tabel.index.name = "Slot \\ Hari"
    return tabel

tabel_hs = bangun_tabel_hari_slot(k_terbaik)
tabel_hs.style.set_properties(**{"white-space": "pre-wrap", "text-align": "left", "vertical-align": "top",
                                 "min-width": "170px", "font-size": "11px"})

# ---------- Verifikasi independen jadwal terbaik (cek langsung dari tabel, bukan dari fungsi fitness) ----------
d = df_jadwal.copy()
assert not d.duplicated(["Hari", "Jam", "Dosen"]).any(),  "Ada dosen bentrok"
assert not d.duplicated(["Hari", "Jam", "Ruang"]).any(),  "Ada ruang bentrok"
assert not d.duplicated(["Hari", "Jam", "Kelas"]).any(),  "Ada kelas bentrok"
kap = {v[0]: v[1] for v in RUANG.values()}
assert all(kap[r] >= m for r, m in zip(d["Ruang"], [KELAS[k] for k in d["Kelas"]])), "Kapasitas kurang"
print("Verifikasi jadwal terbaik: tidak ada bentrok dosen/ruang/kelas, kapasitas terpenuhi ✔")

# ---------- Ekspor CSV & Excel ----------
df_jadwal.to_csv("jadwal_terbaik.csv", index=False, encoding="utf-8-sig")
df_runs.round(4).to_csv("hasil_10_run.csv", index=False, encoding="utf-8-sig")

with pd.ExcelWriter("jadwal_terbaik.xlsx", engine="openpyxl") as xw:
    df_jadwal.to_excel(xw, sheet_name="Daftar Jadwal", index=False)
    tabel_hs.to_excel(xw, sheet_name="Hari x Slot")
    df_runs.round(4).to_excel(xw, sheet_name="Hasil 10 Run", index=False)
    ringkasan.to_excel(xw, sheet_name="Ringkasan 10 Run")

    from openpyxl.styles import Alignment, Font, PatternFill
    for ws in xw.book.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
        for col in ws.columns:
            lebar = max(len(str(c.value).split("\n")[0]) if c.value else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, lebar + 2), 42)
    ws = xw.book["Hari x Slot"]
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 110
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
print("Tersimpan: jadwal_terbaik.csv, jadwal_terbaik.xlsx, hasil_10_run.csv")

# ---------- Grafik konvergensi ----------
fig, axs = plt.subplots(1, 2, figsize=(12, 4.3))

# (a) run terbaik
ax = axs[0]
ax.plot(run_terbaik["riwayat_terbaik"], label="Fitness terbaik", linewidth=2)
ax.plot(run_terbaik["riwayat_rata"], label="Fitness rata-rata populasi", linewidth=1.5, alpha=0.8)
if run_terbaik["gen_valid"] is not None:
    ax.axvline(run_terbaik["gen_valid"], color="green", linestyle="--", label=f"Valid pertama (gen {run_terbaik['gen_valid']})")
ax.set_title(f"Konvergensi run terbaik (seed {run_terbaik['seed']})")
ax.set_xlabel("Generasi"); ax.set_ylabel("Fitness"); ax.legend()

# (b) rata-rata 10 run
ax = axs[1]
M_best = np.array([h["riwayat_terbaik"] for h in hasil_runs])
M_avg  = np.array([h["riwayat_rata"] for h in hasil_runs])
x = np.arange(M_best.shape[1])
ax.plot(x, M_best.mean(0), label="Fitness terbaik (rata-rata 10 run)", linewidth=2)
ax.fill_between(x, M_best.mean(0) - M_best.std(0), np.minimum(1, M_best.mean(0) + M_best.std(0)), alpha=0.2)
ax.plot(x, M_avg.mean(0), label="Fitness rata-rata populasi (rata-rata 10 run)", linewidth=1.5)
ax.set_title("Konvergensi rata-rata 10 run (±1 simpangan baku)")
ax.set_xlabel("Generasi"); ax.set_ylabel("Fitness"); ax.legend()

plt.tight_layout(); plt.savefig("grafik_konvergensi.png", dpi=150); plt.show()
